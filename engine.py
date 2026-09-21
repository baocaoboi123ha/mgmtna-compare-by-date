# -*- coding: utf-8 -*-
"""Load bundled 7.0 engine; emit 4.0-style sheets + unit table under Compare."""

from __future__ import annotations

import importlib.util
import sys
from copy import copy
from pathlib import Path

import lxml.etree  # noqa: F401
import openpyxl  # noqa: F401
from openpyxl.styles import Font


def _orig_engine_path() -> Path:
    if getattr(sys, "frozen", False):
        return Path(getattr(sys, "_MEIPASS")) / "engine_orig.pyc"
    return Path(__file__).resolve().parent / "engine_orig.pyc"


def _load_orig():
    path = _orig_engine_path()
    spec = importlib.util.spec_from_file_location("_engine_orig", path)
    if spec is None or spec.loader is None:
        raise ImportError(f"Khong load duoc engine goc: {path}")
    mod = importlib.util.module_from_spec(spec)
    sys.modules["_engine_orig"] = mod
    spec.loader.exec_module(mod)
    return mod


_orig = _load_orig()
_orig.MATCH_PROFILES["v7"]["include_ok"] = False
_orig.MATCH_PROFILES["v7"]["layout"] = "bc_tp"

_unit_ctx = {"rows": None, "convert_name": "", "attach": False}
_orig_compare_workbook = _orig.compare_workbook
_orig_write_day_compare_v3 = _orig.write_day_compare_v3
_orig_write_tonghop_v3 = _orig.write_tonghop_v3
_orig_build_unit_mismatch_rows = _orig.build_unit_mismatch_rows
_orig_find_unit_convert_path = _orig.find_unit_convert_path
_orig_load_unit_convert = _orig.load_unit_convert

_TONGHOP_SRC_COLS = (1, 2, 3, 4, 7, 8, 9, 10, 11, 12, 5, 6)
_TONGHOP_HEADERS = (
    "Ngay",
    "Sheet",
    "Ma BC",
    "Ma TP",
    "BCqty",
    "TPqty",
    "checkqty(TP-BC)",
    "BCamount",
    "TPamount",
    "checkamount(TP-BC)",
    "So ma lech",
    "So ma lech SL",
)


def _unit_path_from_call(args, kwargs):
    if "unit_convert_path" in kwargs:
        return kwargs.get("unit_convert_path")
    if len(args) >= 6:
        return args[5]
    return None


def compare_workbook(*args, **kwargs):
    """So mã tab passes no unit file — do not attach unit tables on Compare."""
    _unit_ctx["rows"] = None
    _unit_ctx["convert_name"] = ""
    _unit_ctx["attach"] = bool(_unit_path_from_call(args, kwargs))
    try:
        return _orig_compare_workbook(*args, **kwargs)
    finally:
        _unit_ctx["rows"] = None
        _unit_ctx["attach"] = False


def find_unit_convert_path(source, explicit=None):
    """Web: never auto-discover a Unit file from disk; only use the uploaded one."""
    if not explicit:
        return None
    path = _orig_find_unit_convert_path(source, explicit)
    if path is not None:
        _unit_ctx["convert_name"] = getattr(path, "name", str(path))
    return path


def load_unit_convert(path):
    if path is not None:
        _unit_ctx["convert_name"] = getattr(path, "name", str(path))
    return _orig_load_unit_convert(path)


def build_unit_mismatch_rows(*args, **kwargs):
    rows = _orig_build_unit_mismatch_rows(*args, **kwargs)
    if _unit_ctx["attach"]:
        _unit_ctx["rows"] = rows
    return rows


def _sheet(wb, name):
    if name in wb.sheetnames:
        return wb[name]
    safe = _orig._safe_sheet_name(name)
    return wb[safe]


def _style_compare_like_v4(ws, day_label, key_label, tolerance, lech_count, items_tp):
    qty_tol = getattr(_orig, "QTY_TOLERANCE", 0.001)
    ws["A1"] = f"BC (Lines) vs TP (Sheet1) — Ngay {day_label}"
    ws["A2"] = (
        f"Khớp {key_label} (BC) vs Ma San Pham (TP) | Chỉ liệt kê mã LỆCH (ẩn dòng OK) | "
        f"check = TP - BC | Tol tiền = {tolerance} | Tol SL = {qty_tol}"
    )
    status = ws.cell(4, 8).value
    if status in ("Trang thai", "Trạng thái", "Status"):
        ws.cell(4, 8).value = "Status"

    total_row = None
    for row in range(int(ws.max_row or 4), 3, -1):
        if ws.cell(row, 1).value in ("TONG", "TOTAL", "Tổng"):
            total_row = row
            break
    if total_row is None:
        return
    cell = ws.cell(total_row, 1, "TOTAL")
    cell.font = Font(bold=True)
    cell = ws.cell(total_row, 8, f"Lech {int(lech_count)}/{int(items_tp)}")
    cell.font = Font(bold=True)


def _style_tonghop_like_v4(ws, tolerance, day_count):
    qty_tol = getattr(_orig, "QTY_TOLERANCE", 0.001)
    ws["A1"] = "Tong hop BC (Lines) vs TP (Sheet1)"
    ws["A2"] = (
        f"check = TP - BC | Tol tien = {tolerance} | Tol SL = {qty_tol} | "
        f"So ngay = {day_count} | Rule = Item Reference No. vs Ma San Pham — tự nhận cột, chỉ hiện mã lệch"
    )

    last_row = int(ws.max_row or 4)
    snapshot = []
    for row in range(4, last_row + 1):
        snapshot.append(
            [
                (
                    ws.cell(row, col).value,
                    ws.cell(row, col).number_format,
                    copy(ws.cell(row, col).font),
                    copy(ws.cell(row, col).fill),
                    copy(ws.cell(row, col).alignment),
                )
                for col in range(1, 13)
            ]
        )

    for offset, cells in enumerate(snapshot):
        row = 4 + offset
        reordered = [cells[src - 1] for src in _TONGHOP_SRC_COLS]
        for col, (value, fmt, font, fill, alignment) in enumerate(reordered, start=1):
            cell = ws.cell(row, col, _TONGHOP_HEADERS[col - 1] if row == 4 else value)
            if row != 4:
                cell.number_format = fmt
            cell.font = font
            cell.fill = fill
            cell.alignment = alignment


def write_day_compare_v3(
    wb,
    sheet_name,
    pivot_bc,
    pivot_tp,
    day_label,
    key_label,
    tolerance,
    include_ok=True,
    **kwargs,
):
    info = _orig_write_day_compare_v3(
        wb,
        sheet_name,
        pivot_bc,
        pivot_tp,
        day_label,
        key_label,
        tolerance,
        include_ok=include_ok,
        **kwargs,
    )
    ws = _sheet(wb, sheet_name)
    lech_count = int((info or {}).get("LechCount") or 0)
    _style_compare_like_v4(ws, day_label, key_label, tolerance, lech_count, len(pivot_tp))

    rows = _unit_ctx.get("rows") if _unit_ctx["attach"] else None
    _unit_ctx["rows"] = None
    if rows is not None:
        start = int(ws.max_row or 1) + 2
        _orig.write_unit_mismatch_table(
            ws,
            start,
            rows,
            key_label,
            _unit_ctx.get("convert_name") or "",
        )
    return info


def write_tonghop_v3(
    wb,
    day_results,
    tolerance,
    match_label,
    unit_rows=None,
    unit_convert_name=None,
    key_label=None,
):
    label = "Item Reference No. vs Ma San Pham — tự nhận cột, chỉ hiện mã lệch"
    _orig_write_tonghop_v3(
        wb,
        day_results,
        tolerance,
        label,
        unit_rows=None,
        unit_convert_name=None,
        key_label=key_label,
    )
    _style_tonghop_like_v4(wb["TongHop"], tolerance, len(day_results))


_orig.find_unit_convert_path = find_unit_convert_path
_orig.load_unit_convert = load_unit_convert
_orig.build_unit_mismatch_rows = build_unit_mismatch_rows
_orig.write_day_compare_v3 = write_day_compare_v3
_orig.write_tonghop_v3 = write_tonghop_v3

for _name in dir(_orig):
    if _name.startswith("_") or _name == "compare_workbook":
        continue
    globals()[_name] = getattr(_orig, _name)

MATCH_PROFILES = _orig.MATCH_PROFILES
compare_workbook = compare_workbook
CompareError = _orig.CompareError
FileLockedError = _orig.FileLockedError
write_day_compare_v3 = write_day_compare_v3
write_tonghop_v3 = write_tonghop_v3
find_unit_convert_path = find_unit_convert_path
load_unit_convert = load_unit_convert
build_unit_mismatch_rows = build_unit_mismatch_rows
