# -*- coding: utf-8 -*-
"""Web 8.0 engine: unit table redesign, table borders, ByDate 8.0 output."""

from __future__ import annotations

import importlib.util
import sys
from copy import copy
from pathlib import Path

# Keep PyInstaller from dropping runtime deps loaded via engine_orig.pyc
import lxml.etree  # noqa: F401
import openpyxl  # noqa: F401
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

_THIN_BORDER = Border(
    left=Side(style="thin", color="808080"),
    right=Side(style="thin", color="808080"),
    top=Side(style="thin", color="808080"),
    bottom=Side(style="thin", color="808080"),
)


def _apply_table_borders(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    """Draw thin borders around every cell in a rectangular table range."""
    if end_row < start_row or end_col < start_col:
        return
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            ws.cell(row, col).border = _THIN_BORDER


def _orig_engine_path() -> Path:
    if getattr(sys, "frozen", False):
        return Path(getattr(sys, "_MEIPASS")) / "engine_orig.pyc"
    return Path(__file__).resolve().parent / "engine_orig.pyc"


def _load_orig():
    path = _orig_engine_path()
    spec = importlib.util.spec_from_file_location("_engine_orig_v8", path)
    if spec is None or spec.loader is None:
        raise ImportError(f"Khong load duoc engine goc: {path}")
    mod = importlib.util.module_from_spec(spec)
    sys.modules["_engine_orig_v8"] = mod
    spec.loader.exec_module(mod)
    return mod


_orig = _load_orig()
_orig.MATCH_PROFILES["v7"]["include_ok"] = False
_orig.MATCH_PROFILES["v7"]["layout"] = "bc_tp"
_orig.MATCH_PROFILES["v7"]["output_suffix"] = " - ByDate 8.0.xlsx"
_orig.MATCH_PROFILES["v7"]["label"] = "Item Reference No. vs Ma san pham"

_unit_ctx = {"rows": None, "convert_name": "", "attach": False}
_orig_compare_workbook = _orig.compare_workbook
_orig_write_day_compare_v3 = _orig.write_day_compare_v3
_orig_write_tonghop_v3 = _orig.write_tonghop_v3
_orig_build_unit_mismatch_rows = _orig.build_unit_mismatch_rows
_orig_write_unit_mismatch_table = _orig.write_unit_mismatch_table
_orig_find_unit_convert_path = _orig.find_unit_convert_path
_orig_load_unit_convert = _orig.load_unit_convert

_TONGHOP_SRC_COLS = (1, 2, 3, 4, 7, 8, 9, 10, 11, 12, 5, 6)
_TONGHOP_HEADERS = (
    "Ngay",
    "Sheet",
    "Ma BC",
    "Ma TP",
    "BCqty",
    "TPqty",
    "checkqty(TP-BC)",
    "BCamount",
    "TPamount",
    "checkamount(TP-BC)",
    "So ma lech",
    "So ma lech SL",
)


def _unit_path_from_call(args, kwargs):
    if "unit_convert_path" in kwargs:
        return kwargs.get("unit_convert_path")
    if len(args) >= 6:
        return args[5]
    return None


def compare_workbook(*args, **kwargs):
    """So mã tab passes no unit file — do not attach unit tables on Compare."""
    _unit_ctx["rows"] = None
    _unit_ctx["convert_name"] = ""
    _unit_ctx["attach"] = bool(_unit_path_from_call(args, kwargs))
    try:
        return _orig_compare_workbook(*args, **kwargs)
    finally:
        _unit_ctx["rows"] = None
        _unit_ctx["attach"] = False


def find_unit_convert_path(source, explicit=None):
    """Web: never auto-discover a Unit file from disk; only use the uploaded one."""
    if not explicit:
        return None
    path = _orig_find_unit_convert_path(source, explicit)
    if path is not None:
        _unit_ctx["convert_name"] = getattr(path, "name", str(path))
    return path


def load_unit_convert(path):
    if path is not None:
        _unit_ctx["convert_name"] = getattr(path, "name", str(path))
    return _orig_load_unit_convert(path)


def _first_factor(raw) -> float | None:
    """Parse the first Qty/Base factor from a stored FactorBC/FactorTP value."""
    if raw is None or raw == "":
        return None
    text = str(raw).split(";")[0].strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _first_unit(raw) -> str:
    if raw is None:
        return ""
    return str(raw).split(";")[0].strip()


def _format_ratio(value: float) -> str:
    if abs(value - round(value)) < 1e-9:
        return str(int(round(value)))
    return _orig.format_unit_factor(value)


def _convert_formula(unit_bc: str, factor_bc: float | None, unit_tp: str, factor_tp: float | None, unit_base: str) -> str:
    """Explain TP↔BC conversion and clarify Unit Base when factor = 1."""
    parts: list[str] = []
    if factor_bc and factor_tp and factor_bc != 0 and unit_bc and unit_tp:
        parts.append(f"1 {unit_tp} = {_format_ratio(factor_tp / factor_bc)} {unit_bc}")
    elif factor_tp and unit_tp and unit_base:
        parts.append(f"1 {unit_tp} = {_format_ratio(factor_tp)} {unit_base}")
    elif factor_bc and unit_bc and unit_base:
        parts.append(f"1 {unit_bc} = {_format_ratio(factor_bc)} {unit_base}")

    # When Unit Base is exactly 1 of BC/TP unit, spell that out to avoid CAI vs CHAI confusion
    if unit_base:
        if factor_bc is not None and abs(factor_bc - 1.0) < 1e-9 and unit_bc:
            parts.append(f"Unit Base {unit_base} = 1 {unit_bc}")
        elif factor_tp is not None and abs(factor_tp - 1.0) < 1e-9 and unit_tp:
            parts.append(f"Unit Base {unit_base} = 1 {unit_tp}")
        elif (
            factor_bc is not None
            and factor_tp is not None
            and abs(factor_bc - 1.0) >= 1e-9
            and abs(factor_tp - 1.0) >= 1e-9
            and unit_bc
            and unit_tp
        ):
            parts.append(f"1 {unit_bc} = {_format_ratio(factor_bc)} {unit_base}")
            parts.append(f"1 {unit_tp} = {_format_ratio(factor_tp)} {unit_base}")

    return " | ".join(parts)


def _enrich_unit_rows_for_display(rows: list, pivot_bc: dict | None = None, pivot_tp: dict | None = None) -> list:
    """Attach sold qty, conversion formula, and Unit-Base qty diff (TP - BC)."""
    pivot_bc = pivot_bc or {}
    pivot_tp = pivot_tp or {}
    for item in rows:
        key = item.get("Key")
        qty_bc = float(_orig._qty_of(pivot_bc, key) or 0) if key is not None else 0.0
        qty_tp = float(_orig._qty_of(pivot_tp, key) or 0) if key is not None else 0.0
        f_bc = _first_factor(item.get("FactorBC"))
        f_tp = _first_factor(item.get("FactorTP"))
        unit_bc = _first_unit(item.get("UnitBC"))
        unit_tp = _first_unit(item.get("UnitTP"))
        unit_base = str(item.get("UnitBase") or "").strip()

        item["QtySoldBC"] = qty_bc
        item["QtySoldTP"] = qty_tp
        item["ConvertFormula"] = _convert_formula(unit_bc, f_bc, unit_tp, f_tp, unit_base)
        if f_bc is None or f_tp is None:
            item["QtyDiffBase"] = ""
        else:
            item["QtyDiffBase"] = qty_tp * f_tp - qty_bc * f_bc
    return rows


def build_unit_mismatch_rows(*args, **kwargs):
    rows = _orig_build_unit_mismatch_rows(*args, **kwargs)
    if _unit_ctx["attach"]:
        _unit_ctx["rows"] = rows
    return rows


def write_unit_mismatch_table(ws, start_row, rows, key_label, convert_name):
    """Unit mismatch table: sold Qty then Unit, conversion formula, checkqty Base."""
    title_row = start_row
    ws.cell(
        title_row,
        1,
        "So sánh Unit bán ra — lệch ở mức Unit Base (cùng hệ số về Unit Base thì bỏ qua, ví dụ Cái-1 = Hộp-1)",
    )
    ws.cell(title_row, 1).font = Font(name="Calibri", size=13, bold=True, color="833C0C")
    ws.cell(
        title_row + 1,
        1,
        f"Nguồn quy đổi: {convert_name}"
        "  |  Lines = Unit of Measure Code  |  Sheet1 = Đơn vị  |  "
        "Qty BC/TP = số lượng bán thực tế  |  "
        "checkqty Base = QtyTP×hệ sốTP − QtyBC×hệ sốBC (về Unit Base)",
    )
    ws.cell(title_row + 1, 1).font = Font(name="Calibri", size=10, color=_orig.CLR_GRAY)

    headers = (
        key_label,
        "Item No.",
        "Qty BC",
        "Unit BC (Lines)",
        "Qty TP",
        "Unit TP (Sheet1)",
        "Unit Base",
        "Công thức quy đổi",
        "checkqty Base (TP-BC)",
        "Status",
    )
    header_row = title_row + 3
    for i, title in enumerate(headers, start=1):
        cell = ws.cell(header_row, i, title)
        _orig._style_header_cell(
            cell,
            _orig.CLR_LECH_HDR,
            11,
            wrap=True,
        )
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[header_row].height = 36

    if not rows:
        cell = ws.cell(header_row + 1, 1, "(Không có mã lệch Unit Base)")
        cell.font = Font(name="Calibri", italic=True, color=_orig.CLR_GRAY)
        _apply_table_borders(ws, header_row, header_row, 1, len(headers))
        return

    fill = _orig._fill(_orig.CLR_LECH_BOTH)
    last_data_row = header_row
    for offset, item in enumerate(rows, start=1):
        row = header_row + offset
        last_data_row = row
        values = (
            item.get("Key", ""),
            item.get("ItemNo", ""),
            item.get("QtySoldBC", ""),
            _first_unit(item.get("UnitBC")),
            item.get("QtySoldTP", ""),
            _first_unit(item.get("UnitTP")),
            item.get("UnitBase", ""),
            item.get("ConvertFormula", ""),
            item.get("QtyDiffBase", ""),
            item.get("Status", ""),
        )
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            cell.fill = fill
            cell.font = Font(
                name="Calibri",
                color=_orig.CLR_RED,
                bold=(col == 10),
            )
            if col in (1, 2):
                cell.number_format = "@"
            elif col in (3, 5, 9) and isinstance(value, (int, float)):
                cell.number_format = "0.###"
            elif col == 8:
                cell.alignment = Alignment(wrap_text=True, vertical="center")

    _apply_table_borders(ws, header_row, last_data_row, 1, len(headers))

    extra_widths = {
        1: 22,
        2: 14,
        3: 12,
        4: 16,
        5: 12,
        6: 16,
        7: 12,
        8: 36,
        9: 20,
        10: 16,
    }
    for idx, width in extra_widths.items():
        letter = get_column_letter(idx)
        current = ws.column_dimensions[letter].width or 10
        ws.column_dimensions[letter].width = max(current, width)


def _sheet(wb, name):
    if name in wb.sheetnames:
        return wb[name]
    safe = _orig._safe_sheet_name(name)
    return wb[safe]


def _style_compare_like_v4(ws, day_label, key_label, tolerance, lech_count, items_tp):
    qty_tol = getattr(_orig, "QTY_TOLERANCE", 0.001)
    ws["A1"] = f"BC (Lines) vs TP (Sheet1) — Ngay {day_label}"
    ws["A2"] = (
        f"Khớp {key_label} (BC) vs Ma San Pham (TP) | Chỉ liệt kê mã LỆCH (ẩn dòng OK) | "
        f"check = TP - BC | Tol tiền = {tolerance} | Tol SL = {qty_tol}"
    )
    status = ws.cell(4, 8).value
    if status in ("Trang thai", "Trạng thái", "Status"):
        ws.cell(4, 8).value = "Status"

    total_row = None
    for row in range(int(ws.max_row or 4), 3, -1):
        if ws.cell(row, 1).value in ("TONG", "TOTAL", "Tổng"):
            total_row = row
            break
    if total_row is None:
        _apply_table_borders(ws, 4, 4, 1, 8)
        return
    cell = ws.cell(total_row, 1, "TOTAL")
    cell.font = Font(bold=True)
    cell = ws.cell(total_row, 8, f"Lech {int(lech_count)}/{int(items_tp)}")
    cell.font = Font(bold=True)
    _apply_table_borders(ws, 4, total_row, 1, 8)


def _style_tonghop_like_v4(ws, tolerance, day_count):
    qty_tol = getattr(_orig, "QTY_TOLERANCE", 0.001)
    ws["A1"] = "Tong hop BC (Lines) vs TP (Sheet1)"
    ws["A2"] = (
        f"check = TP - BC | Tol tien = {tolerance} | Tol SL = {qty_tol} | "
        f"So ngay = {day_count} | Rule = Item Reference No. vs Ma San Pham — tự nhận cột, chỉ hiện mã lệch"
    )

    last_row = int(ws.max_row or 4)
    snapshot = []
    for row in range(4, last_row + 1):
        snapshot.append(
            [
                (
                    ws.cell(row, col).value,
                    ws.cell(row, col).number_format,
                    copy(ws.cell(row, col).font),
                    copy(ws.cell(row, col).fill),
                    copy(ws.cell(row, col).alignment),
                )
                for col in range(1, 13)
            ]
        )

    for offset, cells in enumerate(snapshot):
        row = 4 + offset
        reordered = [cells[src - 1] for src in _TONGHOP_SRC_COLS]
        for col, (value, fmt, font, fill, alignment) in enumerate(reordered, start=1):
            cell = ws.cell(row, col, _TONGHOP_HEADERS[col - 1] if row == 4 else value)
            if row != 4:
                cell.number_format = fmt
            cell.font = font
            cell.fill = fill
            cell.alignment = alignment

    # Border the TongHop data block (header + day rows + TONG)
    end_row = 4
    for row in range(4, last_row + 1):
        if any(ws.cell(row, col).value is not None for col in range(1, 13)):
            end_row = row
    _apply_table_borders(ws, 4, end_row, 1, 12)


def write_day_compare_v3(
    wb,
    sheet_name,
    pivot_bc,
    pivot_tp,
    day_label,
    key_label,
    tolerance,
    include_ok=True,
    **kwargs,
):
    info = _orig_write_day_compare_v3(
        wb,
        sheet_name,
        pivot_bc,
        pivot_tp,
        day_label,
        key_label,
        tolerance,
        include_ok=include_ok,
        **kwargs,
    )
    ws = _sheet(wb, sheet_name)
    lech_count = int((info or {}).get("LechCount") or 0)
    _style_compare_like_v4(ws, day_label, key_label, tolerance, lech_count, len(pivot_tp))

    rows = _unit_ctx.get("rows") if _unit_ctx["attach"] else None
    _unit_ctx["rows"] = None
    if rows is not None:
        rows = _enrich_unit_rows_for_display(rows, pivot_bc, pivot_tp)
        start = int(ws.max_row or 1) + 2
        write_unit_mismatch_table(
            ws,
            start,
            rows,
            key_label,
            _unit_ctx.get("convert_name") or "",
        )
    return info


def write_tonghop_v3(
    wb,
    day_results,
    tolerance,
    match_label,
    unit_rows=None,
    unit_convert_name=None,
    key_label=None,
):
    label = "Item Reference No. vs Ma San Pham — tự nhận cột, chỉ hiện mã lệch"
    _orig_write_tonghop_v3(
        wb,
        day_results,
        tolerance,
        label,
        unit_rows=None,
        unit_convert_name=None,
        key_label=key_label,
    )
    _style_tonghop_like_v4(wb["TongHop"], tolerance, len(day_results))


_orig.find_unit_convert_path = find_unit_convert_path
_orig.load_unit_convert = load_unit_convert
_orig.build_unit_mismatch_rows = build_unit_mismatch_rows
_orig.write_unit_mismatch_table = write_unit_mismatch_table
_orig.write_day_compare_v3 = write_day_compare_v3
_orig.write_tonghop_v3 = write_tonghop_v3

for _name in dir(_orig):
    if _name.startswith("_") or _name == "compare_workbook":
        continue
    globals()[_name] = getattr(_orig, _name)

MATCH_PROFILES = _orig.MATCH_PROFILES
compare_workbook = compare_workbook
CompareError = _orig.CompareError
FileLockedError = _orig.FileLockedError
write_day_compare_v3 = write_day_compare_v3
write_tonghop_v3 = write_tonghop_v3
find_unit_convert_path = find_unit_convert_path
load_unit_convert = load_unit_convert
build_unit_mismatch_rows = build_unit_mismatch_rows
write_unit_mismatch_table = write_unit_mismatch_table
